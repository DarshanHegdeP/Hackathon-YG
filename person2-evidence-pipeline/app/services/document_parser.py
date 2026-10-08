from __future__ import annotations

import csv
import hashlib
import io
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from app.schemas import DocumentMetadata, TableData


class DocumentParsingResult:
    def __init__(
        self,
        extracted_text: str,
        tables: List[TableData],
        page_count: int,
        sha256: str,
        is_success: bool = True,
        error_message: Optional[str] = None,
    ):
        self.extracted_text = extracted_text
        self.tables = tables
        self.page_count = page_count
        self.sha256 = sha256
        self.is_success = is_success
        self.error_message = error_message


def compute_sha256(content: bytes) -> str:
    """Computes SHA-256 hash of object bytes for audit non-repudiation."""
    return hashlib.sha256(content).hexdigest()


def parse_document(
    file_content: bytes,
    file_name: str,
    file_type: Optional[str] = None,
) -> DocumentParsingResult:
    """
    Parses PDF, XLSX, CSV, or plain-text files into normalized text, tables, and metadata.
    Handles corrupt or password-protected documents gracefully without raising unhandled exceptions.
    """
    sha256_hash = compute_sha256(file_content)
    lower_name = file_name.lower()
    inferred_type = (file_type or "").lower()

    try:
        if lower_name.endswith(".pdf") or "pdf" in inferred_type:
            return _parse_pdf(file_content, sha256_hash)
        elif lower_name.endswith(".xlsx") or lower_name.endswith(".xls") or "spreadsheet" in inferred_type or "excel" in inferred_type:
            return _parse_xlsx(file_content, sha256_hash)
        elif lower_name.endswith(".csv") or "csv" in inferred_type:
            return _parse_csv(file_content, file_name, sha256_hash)
        else:
            return _parse_plain_text(file_content, sha256_hash)
    except Exception as exc:
        return DocumentParsingResult(
            extracted_text="",
            tables=[],
            page_count=0,
            sha256=sha256_hash,
            is_success=False,
            error_message=f"Document parsing error: {type(exc).__name__} - {str(exc)}",
        )


def _parse_pdf(file_content: bytes, sha256_hash: str) -> DocumentParsingResult:
    # Try pymupdf first, fallback to pypdf
    try:
        import pymupdf  # type: ignore

        doc = pymupdf.open(stream=file_content, filetype="pdf")
        if doc.is_encrypted:
            return DocumentParsingResult(
                extracted_text="",
                tables=[],
                page_count=0,
                sha256=sha256_hash,
                is_success=False,
                error_message="Document is password protected and cannot be processed.",
            )

        page_texts: List[str] = []
        page_count = len(doc)
        for page_num in range(page_count):
            page = doc.load_page(page_num)
            page_texts.append(page.get_text())
        doc.close()

        clean_text = "\n\n".join(page_texts).strip()
        return DocumentParsingResult(
            extracted_text=clean_text,
            tables=[],
            page_count=page_count,
            sha256=sha256_hash,
            is_success=True,
        )
    except Exception as mupdf_err:
        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(io.BytesIO(file_content))
            if reader.is_encrypted:
                return DocumentParsingResult(
                    extracted_text="",
                    tables=[],
                    page_count=0,
                    sha256=sha256_hash,
                    is_success=False,
                    error_message="Document is password protected and cannot be processed.",
                )
            page_texts = [page.extract_text() or "" for page in reader.pages]
            clean_text = "\n\n".join(page_texts).strip()
            return DocumentParsingResult(
                extracted_text=clean_text,
                tables=[],
                page_count=len(reader.pages),
                sha256=sha256_hash,
                is_success=True,
            )
        except Exception as pypdf_err:
            return DocumentParsingResult(
                extracted_text="",
                tables=[],
                page_count=0,
                sha256=sha256_hash,
                is_success=False,
                error_message=f"Failed to parse PDF document: {str(mupdf_err)} / {str(pypdf_err)}",
            )


def _parse_xlsx(file_content: bytes, sha256_hash: str) -> DocumentParsingResult:
    import openpyxl  # type: ignore

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    except Exception as exc:
        return DocumentParsingResult(
            extracted_text="",
            tables=[],
            page_count=0,
            sha256=sha256_hash,
            is_success=False,
            error_message=f"Failed to open Excel workbook: {str(exc)}",
        )

    tables: List[TableData] = []
    text_parts: List[str] = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows_data = list(sheet.iter_rows(values_only=True))
        if not rows_data:
            continue

        # Find first non-empty row as header
        header_row_idx = None
        for idx, row in enumerate(rows_data):
            if any(cell is not None for cell in row):
                header_row_idx = idx
                break

        if header_row_idx is None:
            continue

        raw_headers = rows_data[header_row_idx]
        headers = [str(h).strip() if h is not None else f"Column_{col_i+1}" for col_i, h in enumerate(raw_headers)]

        table_rows: List[List[Any]] = []
        for row in rows_data[header_row_idx + 1:]:
            if any(cell is not None for cell in row):
                # normalize values for JSON serialization
                clean_row = [
                    str(cell) if not isinstance(cell, (int, float, bool, type(None))) else cell
                    for cell in row
                ]
                table_rows.append(clean_row)

        tables.append(TableData(sheetName=sheet_name, headers=headers, rows=table_rows))
        
        # Add human-readable text representation
        text_parts.append(f"Sheet: {sheet_name}")
        text_parts.append(" | ".join(headers))
        for row in table_rows[:15]:  # include sample text
            text_parts.append(" | ".join(str(c) for c in row))

    wb.close()
    extracted_text = "\n".join(text_parts).strip()
    return DocumentParsingResult(
        extracted_text=extracted_text,
        tables=tables,
        page_count=len(tables),
        sha256=sha256_hash,
        is_success=True,
    )


def _parse_csv(file_content: bytes, file_name: str, sha256_hash: str) -> DocumentParsingResult:
    # Decode text
    try:
        decoded_text = file_content.decode("utf-8")
    except UnicodeDecodeError:
        decoded_text = file_content.decode("latin-1", errors="replace")

    reader = csv.reader(io.StringIO(decoded_text))
    rows = list(reader)

    tables: List[TableData] = []
    sheet_name = Path(file_name).stem or "Sheet1"

    if rows:
        headers = [h.strip() for h in rows[0]]
        data_rows = rows[1:]
        tables.append(TableData(sheetName=sheet_name, headers=headers, rows=data_rows))

    return DocumentParsingResult(
        extracted_text=decoded_text.strip(),
        tables=tables,
        page_count=1,
        sha256=sha256_hash,
        is_success=True,
    )


def _parse_plain_text(file_content: bytes, sha256_hash: str) -> DocumentParsingResult:
    try:
        text = file_content.decode("utf-8")
    except UnicodeDecodeError:
        text = file_content.decode("latin-1", errors="replace")

    return DocumentParsingResult(
        extracted_text=text.strip(),
        tables=[],
        page_count=1,
        sha256=sha256_hash,
        is_success=True,
    )
